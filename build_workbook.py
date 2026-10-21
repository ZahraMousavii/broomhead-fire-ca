#!/usr/bin/env python3
"""Write Broomhead_CA_scenarios_and_results.xlsx (scenario inputs, ignition lines, CA results, accuracy, data checks).

    python build_workbook.py --results ../results --out ../results
"""
import argparse, glob, os, sys
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bhca.landscape import find_project_root, load_landscape
from bhca.scenarios import load_scenarios, FAMILIES
from bhca.evaluate import discover_cases, family_ignition, family_resolved
from bhca.fzp import load_fzp

HEAD = PatternFill("solid", fgColor="1F3864")
WARN = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BBBBBB")
COMPASS = {0: "N", 45: "NE", 90: "E", 135: "SE", 180: "S", 225: "SW", 270: "W", 315: "NW"}


def put(ws, df, start_row=1, widths=None, wrap=()):
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start_row, j, col)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (float, np.floating)):
                v = None if not np.isfinite(v) else round(float(v), 3)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            c = ws.cell(i, j, v)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if df.columns[j - 1] in wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col) or min(max(len(str(col)), *(len(str(v)) for v in df[col].head(60))) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 3)
    ws.row_dimensions[start_row].height = 34


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--root", default=None)
    a = ap.parse_args()
    root = a.root or find_project_root(os.path.dirname(os.path.abspath(__file__)))
    scen = load_scenarios(root)
    land = load_landscape(root)
    cases, excluded = discover_cases(root, scen)
    res = pd.read_csv(os.path.join(a.results, "all_scenarios_results.csv"))
    ev = res[res.farsite_run != "none"].copy()
    ev["usable"] = ev.farsite_run == "consistent"

    wb = Workbook()
    ws = wb.active; ws.title = "Read me"
    lines = [
        ("Broomhead cellular-automaton fire spread - scenario inputs and results", True),
        ("", False),
        ("Scenario inputs: wind speed/direction from the run matrix; temperature and RH from the .wxs weather files and fuel moisture from the .fms files in 'weather and fuel moisture'.", False),
        ("North ignition line = output/North ignition line/new. Patch-design runs (P01-P18) are not part of the CA.", False),
        ("", False),
        ("Sheets", True),
        ("Scenario inputs      - what every scenario W01-W28 feeds the CA (and which file each value comes from).", False),
        ("Ignition lines       - the six ignition geometries, their shapefile and the FARSITE folder they correspond to.", False),
        ("CA results           - the CA run for every scenario x ignition line (168 simulations), with the run-matrix arrival time and FARSITE / CA times at the reference cell.", False),
        ("Accuracy vs FARSITE  - per-run agreement of the CA with the FARSITE grids (relative arrival error, correlation, burned-area IoU).", False),
        ("Accuracy summary     - the same, averaged by split, ignition line and scenario group.", False),
        ("Data checks          - problems found in the input data and FARSITE runs.", False),
        ("", False),
        ("Column notes", True),
        ("matrix_arrival_min = arrival time listed in the run matrix = time of the last cell burned in the FARSITE run.", False),
        ("Reference cell = the last cell FARSITE burned among the cells the CA can also burn. ref_cell_farsite_min / ref_cell_ca_min = the two models' arrival times there; ca_vs_farsite_pct_at_ref_cell = 100 x (CA - FARSITE) / FARSITE.", False),
        ("farsite_run: 'consistent' = the FARSITE run used exactly this scenario's inputs; 'inconsistent with scenario' = it did not (see Data checks); 'none' = no FARSITE run exists.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        ws.cell(i, 1, t).font = Font(bold=b, size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 150

    rows = []
    for sid, s in scen.items():
        m = s.moisture[35]
        rows.append({"Scenario": sid, "Sub-test": s.phase, "Wind from": COMPASS.get(int(s.wind_from), ""), "Wind from (deg)": s.wind_from,
                     "Wind speed (km/h)": s.wind_kmh, "Fuel-moisture class": s.class_label, "Temperature (C)": s.temp_c, "RH (%)": s.rh,
                     "1-h (%)": m[0], "10-h (%)": m[1], "100-h (%)": m[2], "Live herb (%)": m[3], "Live woody (%)": m[4],
                     "Moisture file (.fms)": s.fms_file, "Weather file (.wxs)": s.wxs_file, "Wind also stored in this .wxs": "yes" if s.wxs_wind_matches else "no (wind from run matrix)",
                     "Run-matrix dead moisture text": s.matrix_dead_text, "Run-matrix live moisture text": s.matrix_live_text,
                     "Differs from run-matrix text": "yes" if [float(x) for x in s.matrix_dead_text.replace("/", " ").split()][:3] != [m[0], m[1], m[2]] and s.moisture_class == "wet" else ""})
    ws = wb.create_sheet("Scenario inputs")
    put(ws, pd.DataFrame(rows), widths={"Sub-test": 26, "Fuel-moisture class": 34, "Moisture file (.fms)": 38, "Weather file (.wxs)": 42, "Wind also stored in this .wxs": 26})
    # highlight wet rows
    for r in range(2, 30):
        if ws.cell(r, 19).value == "yes":
            for c in range(1, 20):
                ws.cell(r, c).fill = WARN

    frows = []
    for fam in FAMILIES:
        used = [c for c in cases if c.family["key"] == fam["key"]]
        if not family_resolved(fam):
            frows.append({"Ignition line": fam["label"], "FARSITE folder (output/...)": "n/a", "Ignition shapefile used by the CA": "not resolved yet",
                          "Ignition cells (1 m)": None, "FARSITE runs found (W01-W28)": 0, "Consistent with scenario": 0})
            continue
        n = int(family_ignition(root, fam, land.geo).sum())
        source = fam["ignition"] or "derived from FARSITE's own ignition (t=0) cells - no shapefile in the project"
        frows.append({"Ignition line": fam["label"], "FARSITE folder (output/...)": fam["folder"], "Ignition shapefile used by the CA": source,
                      "Ignition cells (1 m)": n, "FARSITE runs found (W01-W28)": len(used), "Consistent with scenario": sum(c.usable for c in used)})
    ws = wb.create_sheet("Ignition lines")
    put(ws, pd.DataFrame(frows), widths={"Ignition line": 26, "FARSITE folder (output/...)": 34, "Ignition shapefile used by the CA": 36})

    keep = ["ignition_line", "scenario", "sub_test", "wind_from_deg", "wind_kmh", "moisture_class", "temp_c", "rh_pct", "fm_1h", "fm_10h", "fm_100h", "fm_live_herb", "fm_live_woody",
            "matrix_arrival_min", "ref_cell_farsite_min", "ref_cell_ca_min", "ca_vs_farsite_pct_at_ref_cell", "ca_full_burn_min", "farsite_run", "rel_mae", "corr", "iou_100", "note"]
    ws = wb.create_sheet("CA results")
    put(ws, res[[c for c in keep if c in res.columns]], widths={"sub_test": 26, "note": 60, "ignition_line": 24}, wrap=("note",))

    ev2 = ev[["ignition_line", "scenario", "split", "farsite_run", "wind_kmh", "wind_from_deg", "moisture_class", "matrix_arrival_min", "farsite_last_arrival_min",
              "ref_cell_farsite_min", "ref_cell_ca_min", "ca_vs_farsite_pct_at_ref_cell", "mae", "rel_mae", "bias", "corr", "iou_50", "iou_100", "note"]]
    ws = wb.create_sheet("Accuracy vs FARSITE")
    put(ws, ev2, widths={"note": 60, "ignition_line": 24, "split": 18}, wrap=("note",))
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 4).value != "consistent":
            for c in range(1, 20):
                ws.cell(r, c).fill = BAD

    u = ev[ev.usable]
    def agg(g):
        return (u.groupby(g)[["rel_mae", "corr", "iou_50", "iou_100", "mae", "bias", "ca_vs_farsite_pct_at_ref_cell"]]
                .agg(lambda x: x.abs().mean() if x.name == "ca_vs_farsite_pct_at_ref_cell" else x.mean()).round(3)
                .rename(columns={"ca_vs_farsite_pct_at_ref_cell": "abs_err_pct_at_ref_cell"}).assign(n_runs=u.groupby(g).size()).reset_index())
    group = {}
    for s in scen.values():
        group[s.id] = s.phase
    u = u.assign(scenario_group=u.scenario.map(group))
    ws = wb.create_sheet("Accuracy summary")
    r0 = 1
    for title, g in [("By split", "split"), ("By ignition line", "ignition_line"), ("By scenario group", "scenario_group"), ("By fuel-moisture class", "moisture_class")]:
        ws.cell(r0, 1, title).font = Font(bold=True, size=12)
        df = agg(g); put(ws, df, start_row=r0 + 1); r0 += len(df) + 4
    ws.freeze_panes = None
    ws.column_dimensions["A"].width = 34

    ws = wb.create_sheet("Data checks")
    checks = []
    dc = {}
    for folder, key in [(f["folder"], f["key"]) for f in FAMILIES] + [("North ignition line", "north")]:
        n = match = 0
        bad = []
        for p in sorted(glob.glob(os.path.join(root, "output", folder, "W??.fzp"))):
            sid = os.path.basename(p)[:3]
            try:
                r = load_fzp(p)
            except Exception:
                continue
            n += 1
            m = scen[sid].reference_times.get(key, np.nan)
            te = float(np.nanmax(r.arrival))
            if abs(te - m) < 0.05:
                match += 1
            else:
                bad.append((sid, round(te, 2), m))
        dc[folder + (" (older files, not used)" if folder == "North ignition line" else "")] = dict(n=n, match=match, bad=bad)
    for k, v in dc.items():
        checks.append({"Check": "FARSITE last-arrival time equals the run-matrix value", "Where": k, "Result": f"{v['match']} of {v['n']} runs match",
                       "Detail": "; ".join(f"{a}: FARSITE {b} vs matrix {c}" for a, b, c in v["bad"][:6])})
    for c in cases:
        if c.issues:
            checks.append({"Check": "FARSITE run inputs differ from the scenario definition (run excluded from scoring)", "Where": c.name, "Result": "flagged", "Detail": "; ".join(c.issues)})
    n_shift = [c for c in cases if c.notes]
    if n_shift:
        checks.append({"Check": "Ignition offset", "Where": ", ".join(sorted({c.family['label'] for c in n_shift})), "Result": f"{len(n_shift)} runs",
                       "Detail": n_shift[0].notes[0]})
    for p, r in excluded:
        checks.append({"Check": "Unreadable FARSITE package", "Where": p, "Result": "skipped", "Detail": r})
    checks.append({"Check": "Fuel-moisture values", "Where": "Run matrix vs .fms", "Result": "wet scenarios differ",
                   "Detail": "Run matrix lists 10-h/100-h of the wet class as 33/28 and 36/26 (W17) and 33/36 (W21); the .fms file (used) has 30/28/26. Live woody of the drought class is 55 in the FlamMap reference table and 60 in the .fms file (used)."})
    checks.append({"Check": "Weather files", "Where": "weather and fuel moisture/weather", "Result": "7 files for 28 scenarios",
                   "Detail": "The .wxs files are the last-saved copies of files edited between FARSITE runs. Each scenario uses the .wxs whose temperature and RH match its moisture class; wind speed and direction are the run-matrix values."})
    put(ws, pd.DataFrame(checks), widths={"Check": 52, "Where": 34, "Result": 26, "Detail": 90}, wrap=("Check", "Detail"))
    out = os.path.join(a.out, "Broomhead_CA_scenarios_and_results.xlsx")
    wb.save(out)
    print("written", out)


if __name__ == "__main__":
    main()
